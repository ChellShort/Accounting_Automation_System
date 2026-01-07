import openpyxl

# 1. Load the existing XLSM template file
#    Crucially, set keep_vba=True to preserve macros
template_filename = r'resources\template\CONTROL DE COTIZACIONES 2026.xlsm'
try:
    wb = openpyxl.load_workbook(template_filename, keep_vba=True, rich_text = True)
except FileNotFoundError:
    print(f"Error: The file '{template_filename}' was not found. Please ensure the path is correct.")
    exit()

# 2. Select the specific worksheet you want to work on
#    You can select by name or use the active sheet
sheet_name = 'CONTROL' # Replace with your sheet name if different

if sheet_name.lower() in [x.lower() for x in wb.sheetnames]:
    ws = wb[sheet_name]
else:
    print(f"Error: Sheet '{sheet_name}' not found in the workbook.")
    exit()

# page = wb.sheetnames[0]
# ws = wb[page]

# input(page)

# 3. Write data to specific cells
ws['D2'] = 'Hello, Python!'
ws.cell(row=2, column=4, value='Ricardo Emmanuel Sánchez Martínez') # Using cell coordinates

# 4. Save the modified workbook with a new name (or overwrite the template)
output_filename = 'filled_template_output.xlsm'
wb.save(output_filename)

print(f"Successfully wrote data to '{output_filename}', preserving original macros.")
